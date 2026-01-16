"""
ドキュメント解析サービス
Excel、Word、Googleスプレッドシートからテキストを抽出
"""

import io
import re
from typing import Optional
import pandas as pd
from openpyxl import load_workbook
from docx import Document


class DocumentParser:
    """各種ドキュメントファイルを解析するサービス"""

    @staticmethod
    def parse_excel(file_content: bytes, filename: str = "") -> str:
        """
        Excelファイルを解析してテキストに変換

        Args:
            file_content: ファイルのバイトデータ
            filename: ファイル名（参考用）

        Returns:
            Markdown形式のテーブルテキスト
        """
        try:
            # BytesIOでファイルを読み込む
            file_stream = io.BytesIO(file_content)

            # 全シートを読み込む
            excel_file = pd.ExcelFile(file_stream, engine='openpyxl')

            result_parts = []

            for sheet_name in excel_file.sheet_names:
                df = pd.read_excel(excel_file, sheet_name=sheet_name)

                if df.empty:
                    continue

                # シート名を追加
                result_parts.append(f"## シート: {sheet_name}\n")

                # DataFrameをMarkdownテーブルに変換
                markdown_table = DocumentParser._dataframe_to_markdown(df)
                result_parts.append(markdown_table)
                result_parts.append("\n")

            if not result_parts:
                return "（Excelファイルにデータがありません）"

            return "\n".join(result_parts)

        except Exception as e:
            return f"（Excelファイルの読み込みエラー: {str(e)}）"

    @staticmethod
    def parse_word(file_content: bytes, filename: str = "") -> str:
        """
        Wordファイルを解析してテキストに変換

        Args:
            file_content: ファイルのバイトデータ
            filename: ファイル名（参考用）

        Returns:
            抽出されたテキスト
        """
        try:
            file_stream = io.BytesIO(file_content)
            doc = Document(file_stream)

            result_parts = []

            # 段落を抽出
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    # 見出しスタイルを検出してMarkdown形式に変換
                    if para.style and para.style.name:
                        style_name = para.style.name.lower()
                        if 'heading 1' in style_name or 'title' in style_name:
                            result_parts.append(f"# {text}")
                        elif 'heading 2' in style_name:
                            result_parts.append(f"## {text}")
                        elif 'heading 3' in style_name:
                            result_parts.append(f"### {text}")
                        else:
                            result_parts.append(text)
                    else:
                        result_parts.append(text)

            # テーブルを抽出
            for table in doc.tables:
                result_parts.append("\n")
                table_data = []
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells]
                    table_data.append(row_data)

                if table_data:
                    df = pd.DataFrame(table_data[1:], columns=table_data[0]) if len(table_data) > 1 else pd.DataFrame([table_data[0]])
                    result_parts.append(DocumentParser._dataframe_to_markdown(df))
                result_parts.append("\n")

            if not result_parts:
                return "（Wordファイルにテキストがありません）"

            return "\n".join(result_parts)

        except Exception as e:
            return f"（Wordファイルの読み込みエラー: {str(e)}）"

    @staticmethod
    def parse_csv(file_content: bytes, filename: str = "") -> str:
        """
        CSVファイルを解析してテキストに変換

        Args:
            file_content: ファイルのバイトデータ
            filename: ファイル名（参考用）

        Returns:
            Markdown形式のテーブルテキスト
        """
        try:
            file_stream = io.BytesIO(file_content)

            # 文字コードを自動検出（UTF-8, Shift-JIS, CP932を試行）
            encodings = ['utf-8', 'shift-jis', 'cp932', 'utf-16']
            df = None

            for encoding in encodings:
                try:
                    file_stream.seek(0)
                    df = pd.read_csv(file_stream, encoding=encoding)
                    break
                except:
                    continue

            if df is None or df.empty:
                return "（CSVファイルにデータがありません）"

            return DocumentParser._dataframe_to_markdown(df)

        except Exception as e:
            return f"（CSVファイルの読み込みエラー: {str(e)}）"

    @staticmethod
    def extract_google_sheet_id(url: str) -> Optional[str]:
        """
        GoogleスプレッドシートURLからシートIDを抽出

        Args:
            url: GoogleスプレッドシートのURL

        Returns:
            シートID または None
        """
        # 複数のURLパターンに対応
        patterns = [
            r'/spreadsheets/d/([a-zA-Z0-9-_]+)',
            r'spreadsheets/d/([a-zA-Z0-9-_]+)',
            r'key=([a-zA-Z0-9-_]+)',
        ]

        for pattern in patterns:
            match = re.search(pattern, url)
            if match:
                return match.group(1)

        return None

    @staticmethod
    def parse_google_sheet_public(url: str) -> str:
        """
        公開されているGoogleスプレッドシートをCSVとして読み込み

        Args:
            url: GoogleスプレッドシートのURL

        Returns:
            Markdown形式のテーブルテキスト
        """
        try:
            sheet_id = DocumentParser.extract_google_sheet_id(url)

            if not sheet_id:
                return "（GoogleスプレッドシートのURLが正しくありません）"

            # 公開シートのCSVエクスポートURL
            csv_url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=csv"

            # pandasで直接読み込み
            df = pd.read_csv(csv_url)

            if df.empty:
                return "（Googleスプレッドシートにデータがありません）"

            return DocumentParser._dataframe_to_markdown(df)

        except Exception as e:
            return f"（Googleスプレッドシートの読み込みエラー: {str(e)}。シートが公開設定になっていることを確認してください）"

    @staticmethod
    def _dataframe_to_markdown(df: pd.DataFrame) -> str:
        """
        DataFrameをMarkdownテーブルに変換

        Args:
            df: pandas DataFrame

        Returns:
            Markdown形式のテーブル文字列
        """
        if df.empty:
            return ""

        # NaN値を空文字に置換
        df = df.fillna("")

        # ヘッダー行
        headers = [str(col) for col in df.columns]
        header_row = "| " + " | ".join(headers) + " |"

        # 区切り行
        separator = "| " + " | ".join(["---"] * len(headers)) + " |"

        # データ行
        data_rows = []
        for _, row in df.iterrows():
            row_values = [str(val).replace("|", "\\|").replace("\n", " ") for val in row]
            data_rows.append("| " + " | ".join(row_values) + " |")

        return "\n".join([header_row, separator] + data_rows)

    @staticmethod
    def detect_file_type(filename: str) -> str:
        """
        ファイル名から種類を検出

        Args:
            filename: ファイル名

        Returns:
            ファイル種類 ('excel', 'word', 'csv', 'unknown')
        """
        filename_lower = filename.lower()

        if filename_lower.endswith(('.xlsx', '.xls', '.xlsm')):
            return 'excel'
        elif filename_lower.endswith(('.docx', '.doc')):
            return 'word'
        elif filename_lower.endswith('.csv'):
            return 'csv'
        else:
            return 'unknown'

    @staticmethod
    def parse_document(file_content: bytes, filename: str) -> str:
        """
        ファイル種類を自動判定して解析

        Args:
            file_content: ファイルのバイトデータ
            filename: ファイル名

        Returns:
            解析されたテキスト
        """
        file_type = DocumentParser.detect_file_type(filename)

        if file_type == 'excel':
            return DocumentParser.parse_excel(file_content, filename)
        elif file_type == 'word':
            return DocumentParser.parse_word(file_content, filename)
        elif file_type == 'csv':
            return DocumentParser.parse_csv(file_content, filename)
        else:
            return f"（未対応のファイル形式です: {filename}）"


# インスタンス
document_parser = DocumentParser()
